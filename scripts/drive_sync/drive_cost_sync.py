#!/usr/bin/env python3
"""
Drive Cost Sync — REKAP COST 2026 (Google Sheets) -> Supabase project_cost_summary
====================================================================================
READ-ONLY ke Google Drive: script ini HANYA membaca spreadsheet.
Sumber : "REKAP COST 2026" tab "PROJECT SUMMARY"
Target : Supabase tabel public.project_cost_summary + public.sync_status
         (via RPC sync_write_cost_summary / sync_status_* dengan token rahasia)

Mode:
  --once     : jalankan sync sekali (untuk cron / manual)
  --daemon   : polling tiap N detik; sync bila requested_at > last_synced_at
               (dipakai tombol "🔄 Refresh Sekarang" dari web app)

Token disimpan di /home/ubuntu/.hermes/drive_sync_token.txt (chmod 600).
"""
import json
import os
import re
import sys
import time
import urllib.request
import urllib.error

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TOKEN_FILE = os.environ.get('DRIVE_SYNC_TOKEN_FILE', '/home/ubuntu/.hermes/drive_sync_token.txt')
SB_URL_FILE = os.path.join(BASE_DIR, 'sb_url.txt')
SB_KEY_FILE = os.path.join(BASE_DIR, 'sb_key.txt')
GOOGLE_TOKEN_FILE = '/home/ubuntu/.hermes/google_token.json'
SPREADSHEET_ID = '1A7UYMsp6ocak8VAsTjVYL6qY26CXCccAHDESPxAngGk'
SHEET_NAME = 'PROJECT SUMMARY'

POLL_INTERVAL = int(os.environ.get('DRIVE_SYNC_POLL_SECONDS', '45'))


def read_text(path):
    with open(path, 'r') as f:
        return f.read().strip()


def http_json(url, method='GET', headers=None, body=None, timeout=60):
    data = None
    if body is not None:
        data = json.dumps(body).encode('utf-8')
    req = urllib.request.Request(url, data=data, method=method)
    for k, v in (headers or {}).items():
        req.add_header(k, v)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            raw = resp.read().decode('utf-8')
            return resp.status, (json.loads(raw) if raw.strip() else None)
    except urllib.error.HTTPError as e:
        raw = e.read().decode('utf-8', errors='replace')
        try:
            parsed = json.loads(raw)
        except Exception:
            parsed = raw[:400]
        return e.code, parsed


def google_access_token():
    """Refresh OAuth access token dari google_token.json."""
    tok = json.load(open(GOOGLE_TOKEN_FILE))
    status, data = http_json(
        'https://oauth2.googleapis.com/token',
        method='POST',
        headers={'Content-Type': 'application/x-www-form-urlencoded'},
        body=None,
    )
    # urllib tak mendukung form langsung di http_json; pakai manual di bawah.
    import urllib.parse
    body = urllib.parse.urlencode({
        'client_id': tok['client_id'],
        'client_secret': tok['client_secret'],
        'refresh_token': tok['refresh_token'],
        'grant_type': 'refresh_token',
    }).encode('utf-8')
    req = urllib.request.Request('https://oauth2.googleapis.com/token', data=body, method='POST')
    req.add_header('Content-Type', 'application/x-www-form-urlencoded')
    with urllib.request.urlopen(req, timeout=60) as resp:
        payload = json.loads(resp.read().decode('utf-8'))
    return payload['access_token']


def fetch_revenue(access_token):
    """Baca TOTAL REVENUE per project dari file 'REV PROJECT.xlsx' (READ-ONLY).

    Sumber: file terbaru di Drive bernama 'REV PROJECT.xlsx' (Sheet1: kolom
    PROJECT NO / NAME / TOTAL REVENUE). Return dict {project_no: revenue}.
    """
    import urllib.parse
    # 1) cari file REV PROJECT.xlsx terbaru
    q = urllib.parse.quote("name = 'REV PROJECT.xlsx' and trashed = false")
    order = urllib.parse.quote('modifiedTime desc')
    url = (f'https://www.googleapis.com/drive/v3/files?q={q}'
           f'&fields=files(id,name,modifiedTime)&orderBy={order}&pageSize=1')
    status, data = http_json(url, headers={'Authorization': f'Bearer {access_token}'})
    if status != 200 or not data.get('files'):
        raise RuntimeError(f'REV PROJECT.xlsx tidak ditemukan (HTTP {status})')
    file_id = data['files'][0]['id']

    # 2) export sebagai xlsx? Tidak — file .xlsx di Drive tidak bisa dibaca via
    #    Sheets API. Gunakan export untuk Google Sheet; untuk .xlsx pakai alt=media
    #    lewat openpyxl.
    import urllib.request
    req = urllib.request.Request(
        f'https://www.googleapis.com/drive/v3/files/{file_id}?alt=media',
        headers={'Authorization': f'Bearer {access_token}'})
    with urllib.request.urlopen(req, timeout=90) as resp:
        raw = resp.read()
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb[wb.sheetnames[0]]

    out = {}
    for r in range(2, ws.max_row + 1):
        no = ws.cell(row=r, column=2).value
        rev = ws.cell(row=r, column=4).value
        if no is None or rev is None:
            continue
        key = str(no).strip()
        try:
            out[key] = float(rev)
        except (TypeError, ValueError):
            continue
    return out


def fetch_petty_cash(access_token):
    """Baca PETTY CASH per cost center (READ-ONLY). Return dict {cost_center: nominal}."""
    import urllib.parse
    rng = urllib.parse.quote('PETTY CASH!A1:I2000')
    url = f'https://sheets.googleapis.com/v4/spreadsheets/{SPREADSHEET_ID}/values/{rng}?majorDimension=ROWS'
    status, data = http_json(url, headers={'Authorization': f'Bearer {access_token}'})
    if status != 200:
        raise RuntimeError(f'PETTY CASH HTTP {status}: {str(data)[:200]}')
    values = data.get('values', [])
    if len(values) < 2:
        return {}
    header = [str(h).strip().lower() for h in values[0]]

    def col(name_part):
        for i, h in enumerate(header):
            if name_part in h:
                return i
        return None

    c_nom = col('nominal')
    c_cc = col('cost center')
    out = {}
    for raw in values[1:]:
        def cell(idx):
            if idx is None or idx >= len(raw):
                return ''
            return str(raw[idx]).strip()
        cc = cell(c_cc)
        nom = parse_rupiah(cell(c_nom))
        if not cc or nom is None:
            continue
        # Petty cash per project (mis. "99 Baroni") → ambil angka depannya.
        m = re.match(r'^\s*\[?(\d+)\]?', cc)
        key = m.group(1) if m else cc
        out[key] = out.get(key, 0.0) + nom
    return out


def fetch_project_summary(access_token):
    """Baca tab PROJECT SUMMARY (READ-ONLY). Return list of dict rows."""
    import urllib.parse
    rng = urllib.parse.quote(f'{SHEET_NAME}!A1:I200')
    url = f'https://sheets.googleapis.com/v4/spreadsheets/{SPREADSHEET_ID}/values/{rng}?majorDimension=ROWS'
    status, data = http_json(url, headers={'Authorization': f'Bearer {access_token}'})
    if status != 200:
        raise RuntimeError(f'Sheets API HTTP {status}: {str(data)[:300]}')
    values = data.get('values', [])
    if not values:
        raise RuntimeError('Sheet kosong / tanpa data')

    header = [str(h).strip().lower() for h in values[0]]

    def col(name_part):
        for i, h in enumerate(header):
            if name_part in h:
                return i
        return None

    c_no = col('project no')
    c_name = col('project name')
    c_status = col('project status')
    c_week = col('current week')
    c_cum = col('cumulative')
    c_paid = col('total paid')
    c_out = col('outstanding')
    c_pend = col('pending')
    c_report = col('last report')

    rows = []
    for raw in values[1:]:
        def cell(idx):
            if idx is None or idx >= len(raw):
                return ''
            return str(raw[idx]).strip()
        no = cell(c_no)
        name = cell(c_name)
        if not no and not name:
            continue
        rows.append({
            'project_no': no,
            'project_name': name,
            'project_status': cell(c_status) or None,
            'current_week_expense': parse_rupiah(cell(c_week)),
            'cumulative_cost': parse_rupiah(cell(c_cum)),
            'total_paid': parse_rupiah(cell(c_paid)),
            'outstanding': parse_rupiah(cell(c_out)),
            'pending_verification': parse_rupiah(cell(c_pend)),
            'last_report_date': parse_date(cell(c_report)),
        })
    return rows


def parse_rupiah(s):
    """'Rp 23,901,863' / '0' -> 23901863.0 / 0.0 ; kosong -> None"""
    if s is None:
        return None
    t = str(s).strip()
    if not t:
        return None
    neg = t.startswith('-')
    digits = ''.join(ch for ch in t if ch.isdigit())
    if not digits:
        return None
    val = float(digits)
    return -val if neg else val


def parse_date(s):
    """'' -> None; '2026-10-09' atau '9/10/2026' -> 'YYYY-MM-DD'."""
    if not s:
        return None
    t = str(s).strip()
    if not t:
        return None
    import re
    m = re.match(r'^(\d{4})-(\d{2})-(\d{2})', t)
    if m:
        return f'{m.group(1)}-{m.group(2)}-{m.group(3)}'
    m = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{4})', t)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        return f'{y:04d}-{mo:02d}-{d:02d}'
    return None


def run_sync(trigger):
    """Satu siklus sync. Return (ok, message, rows)."""
    sb_url = read_text(SB_URL_FILE)
    sb_key = read_text(SB_KEY_FILE)
    sync_token = read_text(TOKEN_FILE)
    headers = {'apikey': sb_key, 'Authorization': f'Bearer {sb_key}', 'Content-Type': 'application/json'}

    try:
        access_token = google_access_token()
        rows = fetch_project_summary(access_token)
    except Exception as e:
        msg = f'Gagal baca Google Sheets: {e}'
        http_json(f'{sb_url}/rest/v1/rpc/sync_status_write', method='POST', headers=headers,
                  body={'p_token': sync_token, 'p_status': 'ERROR', 'p_message': msg[:500],
                        'p_rows': 0, 'p_trigger': trigger})
        return False, msg, 0

    # Revenue & Petty Cash (untuk fitur Closing Proyek) — sumber terpisah.
    # Gagal membaca salah satunya TIDAK menggagalkan sync cost utama: kolom
    # revenue/petty_cash hanya dikirim bila sumbernya berhasil dibaca (RPC
    # mempertahankan nilai lama bila key tidak ada).
    rev_ok = pc_ok = False
    try:
        revenue_map = fetch_revenue(access_token)
        rev_ok = True
    except Exception as e:
        print(f'[sync] WARN gagal baca revenue: {e}', flush=True)
        revenue_map = {}
    try:
        petty_map = fetch_petty_cash(access_token)
        pc_ok = True
    except Exception as e:
        print(f'[sync] WARN gagal baca petty cash: {e}', flush=True)
        petty_map = {}

    for row in rows:
        no = str(row.get('project_no') or '').strip()
        if rev_ok:
            row['revenue'] = revenue_map.get(no, 0.0)
        if pc_ok:
            row['petty_cash'] = petty_map.get(no, 0.0)

    status, data = http_json(f'{sb_url}/rest/v1/rpc/sync_write_cost_summary', method='POST',
                             headers=headers, body={'p_token': sync_token, 'p_rows': rows})
    if status != 200 or not data or not data.get('ok'):
        msg = f'Gagal tulis DB: HTTP {status} {str(data)[:300]}'
        http_json(f'{sb_url}/rest/v1/rpc/sync_status_write', method='POST', headers=headers,
                  body={'p_token': sync_token, 'p_status': 'ERROR', 'p_message': msg[:500],
                        'p_rows': 0, 'p_trigger': trigger})
        return False, msg, 0

    n = data.get('rows', len(rows))
    extras = []
    if rev_ok: extras.append(f'{len(revenue_map)} revenue')
    if pc_ok: extras.append(f'{len(petty_map)} petty cash')

    # Sprint 1 — TRIGGER B: Early warning cost proyek >= 65% dari revenue.
    # Dipanggil SETELAH sync sukses; anti-spam & claim ditangani RPC
    # telegram_send_cost_alerts (kirim ulang hanya bila naik >= 2 poin).
    # Kegagalan alert TIDAK menggagalkan sync (hanya dicatat di pesan).
    alert_info = ''
    try:
        st_a, data_a = http_json(f'{sb_url}/rest/v1/rpc/telegram_send_cost_alerts', method='POST',
                                 headers=headers, body={'p_token': sync_token, 'p_min_step': 2})
        if st_a == 200 and data_a and data_a.get('ok'):
            sent = data_a.get('sent') or []
            if sent:
                projs = ', '.join(f"{s.get('project_no')} ({s.get('pct')}%)" for s in sent)
                alert_info = f' ⚠️ {len(sent)} alert early warning 65% terkirim: {projs}.'
                print(f'[sync] alert cost terkirim: {projs}', flush=True)
            elif (data_a.get('skipped') or 0) > 0:
                print(f"[sync] {data_a.get('skipped')} proyek >=65% dilewati (kenaikan < 2 poin / belum berubah)", flush=True)
        else:
            print(f'[sync] WARN alert cost gagal: HTTP {st_a} {str(data_a)[:200]}', flush=True)
    except Exception as e:
        print(f'[sync] WARN alert cost error: {e}', flush=True)

    msg = f'Sinkron sukses: {n} project dibaca dari tab {SHEET_NAME}' + (f' (+ {", ".join(extras)})' if extras else '') + '.' + alert_info
    http_json(f'{sb_url}/rest/v1/rpc/sync_status_write', method='POST', headers=headers,
              body={'p_token': sync_token, 'p_status': 'OK', 'p_message': msg[:500],
                    'p_rows': n, 'p_trigger': trigger})
    return True, msg, n


def get_status():
    sb_url = read_text(SB_URL_FILE)
    sb_key = read_text(SB_KEY_FILE)
    sync_token = read_text(TOKEN_FILE)
    headers = {'apikey': sb_key, 'Authorization': f'Bearer {sb_key}', 'Content-Type': 'application/json'}
    status, data = http_json(f'{sb_url}/rest/v1/rpc/sync_status_read', method='POST',
                             headers=headers, body={'p_token': sync_token})
    if status == 200 and data and data.get('ok'):
        return data.get('status') or {}
    return {}


def parse_ts(s):
    if not s:
        return None
    try:
        from datetime import datetime, timezone
        return datetime.fromisoformat(str(s).replace('Z', '+00:00'))
    except Exception:
        return None


def daemon_loop():
    print(f'[daemon] polling tiap {POLL_INTERVAL}s — Ctrl+C untuk berhenti', flush=True)
    # Sync otomatis harian: jalankan sekali bila data terakhir lebih tua dari
    # AUTO_SYNC_HOURS (default 24 jam). Bisa dimatikan via env DRIVE_SYNC_AUTO=0.
    auto_enabled = os.environ.get('DRIVE_SYNC_AUTO', '1') != '0'
    auto_hours = float(os.environ.get('DRIVE_SYNC_AUTO_HOURS', '24'))
    while True:
        try:
            st = get_status()
            req = parse_ts(st.get('requested_at'))
            last = parse_ts(st.get('last_synced_at'))
            # 1) permintaan manual (tombol Refresh Sekarang)
            if req and (last is None or req > last):
                print(f'[daemon] ada permintaan sync (requested_at={st.get("requested_at")}) — jalan…', flush=True)
                ok, msg, n = run_sync('manual')
                print(f'[daemon] hasil: ok={ok} rows={n} — {msg}', flush=True)
            # 2) sync otomatis harian (bila belum ada data atau sudah kedaluwarsa)
            elif auto_enabled:
                from datetime import datetime, timezone, timedelta
                now = datetime.now(timezone.utc)
                stale = last is None or (now - last) > timedelta(hours=auto_hours)
                if stale:
                    print(f'[daemon] sync otomatis (terakhir: {st.get("last_synced_at") or "belum pernah"}) — jalan…', flush=True)
                    ok, msg, n = run_sync('cron')
                    print(f'[daemon] hasil otomatis: ok={ok} rows={n} — {msg}', flush=True)
        except Exception as e:
            print(f'[daemon] error siklus: {e}', flush=True)
        time.sleep(POLL_INTERVAL)


def main():
    args = sys.argv[1:]
    if '--daemon' in args:
        daemon_loop()
        return
    trigger = 'manual' if '--manual' in args else 'cron'
    ok, msg, n = run_sync(trigger)
    print(json.dumps({'ok': ok, 'rows': n, 'message': msg}, ensure_ascii=False))
    sys.exit(0 if ok else 1)


if __name__ == '__main__':
    main()
